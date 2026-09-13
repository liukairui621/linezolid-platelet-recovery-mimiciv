"""Create editable submission files using the configured bundled Python runtime."""
from pathlib import Path
import re,json,csv,zipfile,shutil,hashlib,datetime
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
R=Path(__file__).resolve().parents[2];O=R/'submission_v0.14';O.mkdir(exist_ok=True)
P='routine_culture_no_recorded_lmwh_shared';S='routine_bacterial_culture72h_shared'
def rows(path,sep=','):
 with (R/path).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f,delimiter=sep))
base=rows('submission_v0.14/disclosure_safe/reports/submission_v0.14/TABLE1_BASELINE_v0.14.csv');ass=rows('reports/CLINICAL_ASSOCIATIONS_v0.11.csv');counts=rows('reports/FIRST_EVENT_COUNTS_v0.11.csv')
source=(O/'manuscript_source.md').read_text(encoding='utf-8');title=source.splitlines()[0][2:]
AUTHORS=['Kairui Liu','Rong Chen','Yipei Huang','Youxing Huang'];AFFIL='Department of Abdominal Surgery, The Second Affiliated Hospital of Guangzhou University of Chinese Medicine, Guangzhou, Guangdong Province, P.R. China'
ADDRESS='No. 111 Dade Road, Yuexiu District, Guangzhou 510000, Guangdong Province, P.R. China'
EMAIL='waiqike7@163.com';PHONE='+86-20-39318791; +86-13632255441'
fund=source.split('### Funding\n\n')[1].split('\n\n')[0]
refs={x['key']:x for x in json.loads((R/'literature/submission_v0.14/verified_references.json').read_text(encoding='utf-8'))['records']}
refs['MIMIC31']=dict(authors=['Johnson A','Bulgarelli L','Pollard T','Gow B','Moody B','Horng S','Celi LA','Mark R'],title='MIMIC-IV (version 3.1)',journal='PhysioNet',year=2024,volume='',pages='',doi='10.13026/kpb9-mt58',source='https://physionet.org/content/mimiciv/3.1/')
refs['REPOSITORY']=dict(authors=['Liu K','Chen R','Huang Y','Huang Y'],title='Linezolid platelet recovery in MIMIC-IV: analysis code and aggregate outputs',url='https://github.com/liukairui621/linezolid-platelet-recovery-mimiciv',accessed='11 September 2026')
order=[]
def cite(m):
 keys=[x.lstrip('@') for x in m.group(1).split(';')]
 for key in keys:
  assert key in refs,key
  if key not in order:order.append(key)
 return '['+', '.join(str(order.index(k)+1) for k in keys)+']'
resolved=re.sub(r'\[(@[^\]]+)\]',cite,source)
def reftext(x):
 au=[a.replace('FIRTH D','Firth D') for a in x.get('authors',[])];name=', '.join(au[:6])+(', et al.' if len(au)>6 else '.')
 if x.get('url'):return f"{name} {x['title']}. {x['url']}. Accessed {x['accessed']}."
 page=x.get('pages','');page=re.sub(r'^(e\d+)-\1$',r'\1',str(page));vol=str(x.get('volume',''))
 journal=x.get('journal','');journal={'Nucleic Acids Research':'Nucleic Acids Res','Journal of the American Statistical Association':'J Am Stat Assoc','Nat. Health.':'Nat Health'}.get(journal,journal)
 return f"{name} {x['title'].rstrip('.')}. {journal}. {x.get('year','')}{';'+vol if vol else ''}{':'+page if page else ''}. https://doi.org/{x['doi']}."
biblio='\n\n'.join(f'{i+1}. '+reftext(refs[k]) for i,k in enumerate(order));resolved=resolved.replace('{{REFERENCES}}',biblio)
(O/'Main_Manuscript.md').write_text(resolved,encoding='utf-8')
(O/'references_numbered.json').write_text(json.dumps([dict(number=i+1,**refs[k]) for i,k in enumerate(order)],ensure_ascii=False,indent=2),encoding='utf-8')
(O/'references.ris').write_text('\n'.join('TY  - '+('ELEC' if k=='REPOSITORY' else 'DATA' if k=='MIMIC31' else 'JOUR')+'\nTI  - '+refs[k]['title']+'\n'+''.join('AU  - '+a+'\n' for a in refs[k].get('authors',[]))+'PY  - '+str(refs[k].get('year','2026'))+'\nJO  - '+refs[k].get('journal','GitHub')+'\nDO  - '+refs[k].get('doi','')+'\nUR  - '+refs[k].get('url',refs[k].get('source',''))+'\nER  -\n' for k in order),encoding='utf-8')
def normal_doc(double=False,linenums=False):
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.85);sec.left_margin=sec.right_margin=Inches(1)
 for nm in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
  st=d.styles[nm];st.font.name='Times New Roman';st.font.size=Pt(12);st.font.color.rgb=RGBColor(0,0,0)
  rf=st._element.get_or_add_rPr().get_or_add_rFonts();rf.set(qn('w:eastAsia'),'Microsoft YaHei')
  for key in list(rf.attrib):
   if 'theme' in key.lower():del rf.attrib[key]
  for border in st._element.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
  st.paragraph_format.space_before=Pt(0 if nm=='Normal' else 10);st.paragraph_format.space_after=Pt(6)
  if nm!='Normal':st.font.bold=True;st.paragraph_format.keep_with_next=True
 d.styles['Title'].font.size=Pt(14)
 d.styles['Normal'].paragraph_format.line_spacing=2 if double else 1.15
 d.styles['Normal'].paragraph_format.widow_control=True
 if linenums:
  el=OxmlElement('w:lnNumType');el.set(qn('w:countBy'),'1');el.set(qn('w:restart'),'continuous');el.set(qn('w:distance'),'300');sec._sectPr.append(el)
 footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 ru=footer.add_run('Page ');ru.font.size=Pt(9);fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
 return d
def para(d,text,style=None):
 p=d.add_paragraph(text,style);return p
def md(d,text,with_authors=False):
 for block in text.split('\n\n'):
  block=block.strip()
  if not block:continue
  if block.startswith('# '):
   para(d,block[2:],'Title')
   if with_authors:
    para(d,', '.join(AUTHORS));para(d,AFFIL);para(d,'Correspondence: Youxing Huang, MD. '+EMAIL)
  elif block.startswith('### '):para(d,block[4:],'Heading 2')
  elif block.startswith('## '):para(d,block[3:],'Heading 1')
  else:para(d,block)
def table(d,title,headers,data,widths=None,footnote='',font=9):
 cap=para(d,title,'Heading 2');cap.paragraph_format.keep_with_next=True
 t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths is None:widths=[6.5/len(headers)]*len(headers)
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for j,h in enumerate(headers):t.rows[0].cells[j].text=h
 for row in data:
  cells=t.add_row().cells
  for j,v in enumerate(row):cells[j].text=re.sub(r'(?<!\w)-(?=\d)', '−', str(v))
 pr=t._tbl.tblPr;old=pr.find(qn('w:tblBorders'))
 if old is not None:pr.remove(old)
 borders=OxmlElement('w:tblBorders')
 for name in ['top','left','bottom','right','insideH','insideV']:
  el=OxmlElement('w:'+name);el.set(qn('w:val'),'single' if name in ['top','bottom'] else 'nil');el.set(qn('w:sz'),'8');el.set(qn('w:color'),'000000');borders.append(el)
 pr.append(borders)
 for i,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if i==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
  for j,cell in enumerate(row.cells):
   cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for side,value in [('top','65'),('bottom','65'),('left','60'),('right','60')]:el=OxmlElement('w:'+side);el.set(qn('w:w'),value);el.set(qn('w:type'),'dxa');mar.append(el)
   cp.append(mar)
   if i==0:
    cb=OxmlElement('w:tcBorders')
    for side in ['top','bottom']:el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'8' if side=='top' else '5');el.set(qn('w:color'),'000000');cb.append(el)
    cp.append(cb)
   for p in cell.paragraphs:
    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.1;p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
    if len(data)<=18:p.paragraph_format.keep_with_next=(i<len(data) or bool(footnote))
    for run in p.runs:run.font.size=Pt(font);run.font.bold=i==0
 if footnote:
  p=para(d,footnote)
  p.paragraph_format.line_spacing=1.1
  p.paragraph_format.keep_together=True
  for ru in p.runs:ru.font.size=Pt(9)
 return t
def fnum(x,dec=3):
 if x=='Suppressed':return 'Suppressed'
 try:return f'{float(x):.{dec}f}' if x and x!='NA' else '—'
 except:return '—'
def bdata(pop,selected=None):
 z=[x for x in base if x['population']==pop and (selected is None or x['key'] in selected)]
 return [[x['label'].replace('10^9','10⁹').replace('Cardiac Vascular Intensive Care Unit (CVICU)','CVICU').replace('Coronary Care Unit (CCU)','CCU').replace('Medical Intensive Care Unit (MICU)','MICU').replace('Medical/Surgical Intensive Care Unit (MICU/SICU)','MICU/SICU').replace('Surgical Intensive Care Unit (SICU)','SICU').replace('Trauma SICU (TSICU)','TSICU'),x['LZD_before'],x['VAN_before'],x['LZD_after'],x['VAN_after'],fnum(x['abs_SMD_before']),fnum(x['abs_SMD_after'])] for x in z]
BH=['Characteristic','LZD\nunweighted','VAN\nunweighted','LZD\nweighted','VAN\nweighted','|SMD|\nbefore','|SMD|\nafter'];BW=[1.9,.78,.82,.78,.82,.7,.7]
BF='Continuous values are observed mean (SD); missing values are not filled in this table. Unweighted binary values are n (%), weighted values are %. Weighted SD uses normalized weighted population variance. SMDs use the fitted model scale, including transformations and median filling. Suppressed denotes small cells or linked values removed to prevent count reconstruction, including SMDs and categorical totals; — denotes a constant or absent term. LZD, linezolid; VAN, vancomycin.'
def astat(pop,metric):return next(x for x in ass if x['population']==pop and x['metric']==metric)
def evnum(pop,metric,drug):
 if '__' not in metric:return '—'
 ep,st=metric.split('__');z=[x for x in counts if x['population']==pop and x['drug']==drug and x['endpoint']==ep and x['state']==st];return z[0]['n'] if len(z)==1 else '—'
mainmetrics=[('hospital__recovery','Hospital confirmed recovery'),('hospital__death','Death before confirmation'),('hospital__live_exit','Live discharge before confirmation'),('hospital_single__recovery','Hospital single threshold'),('hospital72__recovery','Confirmation gap 24–72 h'),('hospital_timed__recovery','Hospital timed-transfusion screened'),('physical__recovery','Physical ICU confirmed recovery')]
def adata(pop,metrics):
 out=[]
 for metric,label in metrics:
  x=astat(pop,metric);out.append([label,evnum(pop,metric,'linezolid')+' / '+evnum(pop,metric,'vancomycin'),f'{100*float(x["LZD"]):.2f}',f'{100*float(x["VAN"]):.2f}',f'{100*float(x["RD"]):.2f} ({100*float(x["lower"]):.2f}, {100*float(x["upper"]):.2f})'])
 return out
# Main manuscript, all text editable; tables follow references without manual breaks.
d=normal_doc(True,True);md(d,resolved,True)
selected=['age','female','known_platelets','known_creatinine','baseline_inr','baseline_bilirubin','days_since_admission','rrt_documented_prior24h','invasive_vent_documented_prior24h','pressors_documented_prior6h','prior_recorded_marrow','respiratory72h','nonscreen_blood72h','nonscreen_urine72h','observed_creatinine_aki','beta_lactam_available24h']
table(d,'Table 1. Selected baseline characteristics before and after overlap weighting',BH,bdata(P,selected),BW,'Admissions: LZD 65; VAN 2492. '+BF+' Omitted terms include larger pre-weighting imbalances; the complete term list and missingness are in Additional file 1, Tables S2–S5.')
table(d,'Table 2. Recovery and competing-event probabilities',['Outcome','Events\nLZD / VAN','LZD\nweighted %','VAN\nweighted %','RD (95% CI), pp'],adata(P,mainmetrics)+adata(S,[('hospital__recovery','Full routine-culture cohort confirmed recovery')]),[2.15,.83,.74,.74,2.04],'Denominators: LZD 65 / VAN 2492 for the first seven rows; LZD 65 / VAN 2528 for the full routine-culture sensitivity row. Events are unweighted; probabilities and intervals use overlap weighting and 2000 patient-cluster refits. Only the first row is primary. Small event counts are suppressed. Physical ICU exit changes the estimand. Transfusion-record-screened recovery is not spontaneous recovery. Additional frozen contrasts are in Additional file 1 and the machine-readable companion.')
d.save(O/'Main_Manuscript.docx')
# Standalone title page.
d=normal_doc();para(d,title,'Title');para(d,', '.join(AUTHORS));para(d,AFFIL)
para(d,'Corresponding author','Heading 1');para(d,'Youxing Huang, MD\nDepartment of Abdominal Surgery\nThe Second Affiliated Hospital of Guangzhou University of Chinese Medicine\n'+ADDRESS+'\nTel: '+PHONE+'\nEmail: '+EMAIL)
para(d,'Article type and target','Heading 1');para(d,'Research article — BMC Pharmacology and Toxicology');para(d,'Running title: Antibiotic initiation and platelet recovery')
para(d,'Funding','Heading 1');para(d,fund);para(d,'Competing interests','Heading 1');para(d,'The authors declare that they have no competing interests.');d.save(O/'Title_Page.docx')
# Cover letter: no unconfirmed author-approval or exclusivity assertions.
d=normal_doc();para(d,'12 September 2026');para(d,'Editorial Office\nBMC Pharmacology and Toxicology');para(d,'Dear Editor,')
para(d,'Please consider our Research article, “'+title+',” for publication in BMC Pharmacology and Toxicology.')
para(d,'This study examines documented recovery in critically ill adults who already had thrombocytopenia when initial linezolid or vancomycin treatment was recorded. It complements incident-thrombocytopenia studies by comparing confirmed recovery at a defined treatment decision point while retaining early death and live discharge as competing events.')
para(d,'The selected MIMIC-IV cohort included 65 linezolid and 2492 vancomycin initiations. Fourteen and 1019 admissions met the primary recovery definition. Overlap-weighted probabilities were 21.22% and 32.13% (difference, −10.92 percentage points; 95% CI, −20.12 to −3.19). Route-restricted and multiple-imputation sensitivity analyses produced similar estimates. The study contributes an active-treatment comparison for patients who already have thrombocytopenia when antibiotic treatment begins.')
para(d,'The question fits the journal’s clinical pharmacology and pharmacoepidemiology scope. The submission includes editable tables, final-size composite figures, analysis definitions, aggregate source outputs, and a reproducibility repository. Individual MIMIC records remain subject to credentialed access. Funding and the absence of competing interests are disclosed in the manuscript.')
para(d,'[BEFORE SENDING: the corresponding author must confirm approval by every author, originality and absence of simultaneous submission, and the applicable ethics/access declarations. These statements have not been supplied and must not be assumed.]')
para(d,'Sincerely,\nYouxing Huang, MD\nCorresponding author\n'+AFFIL+'\n'+EMAIL);d.save(O/'Cover_Letter.docx')
# Supplementary methods and complete tables.
d=normal_doc();para(d,'Additional file 1. Supplementary methods and results','Title');para(d,title)
para(d,'S1. Preserved design and source definitions','Heading 1')
para(d,'The primary and full routine-culture populations correspond to the frozen v0.11 definitions. The local analysis plan was frozen on 11 September 2026 after earlier analyses of overlapping data. No analyses in this supplement create a new confirmatory family. Three of five baseline candidates passed numerical and balance criteria; population choice also incorporated the target infection-evaluation context and measured LMWH support. No candidate was selected on a newly estimated recovery contrast.')
table(d,'Table S1. Baseline candidate populations and readiness',['Population','LZD / VAN','Maximum |SMD|','Gate passed'],[[x['population'].replace('_',' '),x['n_LZD']+' / '+x['n_VAN'],fnum(x['max_SMD']),x['pass']] for x in rows('reports/submission_v0.14/READINESS_SELECTION_v0.12.csv')],[3.65,.8,1.15,.6],font=9)
para(d,'Baseline laboratory information required specimen time strictly before initiation and storage time no later than initiation. Platelets used the latest unique numeric result in the preceding 24 hours (item IDs 51265 or 53189; K/uL). Prior platelet change was current baseline minus the latest unique count between 72 and 24 hours before initiation. Creatinine used 50912/52546 in mg/dL; international normalized ratio used 51237; bilirubin used 50885 in mg/dL, each within 24 hours. Observed creatinine AKI used available prior minima and either an increase of at least 0.3 mg/dL over 48 hours or a ratio of at least 1.5 within seven days; it is a recorded creatinine proxy, not adjudicated complete KDIGO AKI.')
para(d,'Recorded renal replacement and invasive ventilation used the preceding 24 hours; vasopressor support used the preceding six hours. Infection collection sources used 72 hours, available organism and selected heparin/marrow-related medication records seven days, and intravenous beta-lactams 24 hours. Baseline eMAR availability and calendar era addressed source coverage imperfectly. Original unit, route, drug, organism and procedure code definitions are preserved in configuration files and scripts 71, 73, 74 and their audited dependencies. Medication name matching represents a recorded-exposure proxy and not an exhaustive pharmacologic classification.')
tax=rows('config/microbiology_test_taxonomy_v0.10.tsv','\t');names=[x['test_name'] for x in tax if x['routine_bacterial'].lower()=='true']
para(d,'The exact routine-culture test names were: '+ '; '.join(names)+'. Specimen screening/surveillance labels overrode a routine test name. Date-only sample/result records used next-midnight availability; actual results entered baseline organism flags only when available by initiation. These rules are conservative classification choices and need not capture every clinically relevant culture.')
para(d,'Physical ICU observation merged overlapping or adjacent intervals for CVICU, CCU, MICU, MICU/SICU, SICU and TSICU. A positive non-ICU gap ended follow-up for that window. A separate historical ICU-stay-record endpoint is retained below to distinguish administrative stay IDs from uninterrupted physical location. Neither window resumes after its first exit.')
para(d,'A recovery pair required post-initiation high specimens separated by at least 24 hours and no intervening observed low specimen. The 24–72-hour version additionally capped the gap. Timed transfusion screening rejected each candidate pair if a captured input interval overlapped first-high minus 48 hours through confirmation; the dated version added entire procedure days. Patient denominators and weights were unchanged. Invalidated early pairs did not prevent a later pair. The final primary source refresh included 74328 platelet laboratory records, 4646 timed platelet inputs, and 28 dated procedure records across the union cohort. Counts are source records rather than distinct patients or transfusion episodes. The eMAR search found no platelet-name records, so transfusion screening used timed inputs and dated procedure records.')
para(d,'R 4.3.3 with brglm2 1.1.0 fitted mean bias-reduced binomial-logit models, initially allowing 200 iterations and a prespecified fallback of 2000 iterations at a slower step. The baseline missing-value medians were recalculated within every patient draw. Patients were resampled within their observed treatment-pattern strata and all eligible admissions of a sampled patient were retained together. There were 2000 valid draws in each population, with three primary and five full-cohort numerical fallbacks. Percentile intervals used R type 7 quantiles, without multiplicity adjustment.')
para(d,'S2. Baseline summaries and completeness','Heading 1')
table(d,'Table S2. Full primary-cohort baseline summary',BH,bdata(P),BW,'Admissions: LZD 65; VAN 2492. '+BF)
table(d,'Table S3. Full routine-culture cohort baseline summary',BH,bdata(S),BW,'Admissions: LZD 65; VAN 2528. '+BF)
miss=rows('reports/submission_v0.14/TABLE_S1_MISSINGNESS_v0.12.csv')
for pop,num in [(P,4),(S,5)]:
 zz=[x for x in miss if x['population']==pop];dd=[]
 for label in dict.fromkeys(x['label'] for x in zz):
  rr=[x for x in zz if x['label']==label];dd.append([label]+[next(x['missing'] for x in rr if x['drug']==g)+' / '+next(x['total'] for x in rr if x['drug']==g) for g in ['LZD','VAN']])
 table(d,f'Table S{num}. Missing original continuous values: '+('primary' if pop==P else 'full routine-culture'),['Variable','LZD missing / total','VAN missing / total'],dd,[3.5,1.5,1.5],'Zero is explicitly shown. Counts below 10 are suppressed. Missingness pertains to original source values, before median filling.')
para(d,'The machine-readable TABLE_S2_MODEL_BALANCE_v0.12.csv lists model-matrix terms on their transformed scales. Values that permit reconstruction of small binary cells, their associated SMDs and categorical totals are suppressed separately within each candidate population. Complete precision remains in the internal analysis records; the reported maximum absolute SMD uses every fitted term before disclosure suppression. Table S2/S3 displays original units. Constant culture and primary LMWH terms were omitted deterministically, not ignored after seeing an adverse balance result.')
para(d,'S3. Complete frozen clinical contrasts','Heading 1')
prefix={'hospital':'Hospital confirmed','hospital_single':'Hospital single','hospital72':'Hospital 24–72 h','hospital_timed':'Hospital timed-screened','hospital_dated':'Hospital timed+date screened','physical':'Physical ICU confirmed','physical_single':'Physical ICU single','physical_timed':'Physical ICU timed-screened*','physical_dated':'Physical ICU timed+date screened*','icu':'ICU-stay record confirmed','icu_single':'ICU-stay record single','icu_timed':'ICU-stay record timed-screened','icu_dated':'ICU-stay record timed+date screened'}
for pop,num in [(P,6),(S,7)]:
 mm=[(x['metric'],prefix[x['metric'].split('__')[0]]+'; '+{'recovery':'recovery','death':'death before event','live_exit':'live exit before event'}[x['metric'].split('__')[1]]) for x in ass if x['population']==pop and '__' in x['metric']]
 table(d,f'Table S{num}. All first-event contrasts: '+('primary' if pop==P else 'full routine-culture'),['Endpoint / first event','Events\nLZD / VAN','LZD %','VAN %','RD (95% CI), pp'],adata(pop,mm),[2.15,.83,.74,.74,2.04],'Each endpoint separately partitions recovery, competing death, competing live exit and no first event. Raw event counts are unweighted. Only the hospital confirmed primary-cohort recovery contrast was primary. *The primary physical ICU screened interval upper endpoint is −0.08 pp, with Monte Carlo SE 0.175 pp, so its exclusion of zero is not separate positive evidence. Dated-screened estimates coincide with timed-screened estimates here. All intervals are exploratory beyond the primary contrast.',font=9)
for pop,num in [(P,8),(S,9)]:
 dd=[]
 for x in ass:
  if x['population']!=pop or '__' in x['metric']:continue
  metric=x['metric'];scale=100 if metric in ['never_threshold','unconfirmed_observed_low','unconfirmed_insufficient_clock','unconfirmed_no_late_measurement'] else 1;unit='%' if scale==100 else ('days' if 'days' in metric else ('tests/day' if 'per_day' in metric else 'tests'))
  dd.append([metric.replace('_',' '),unit,f'{scale*float(x["LZD"]):.3f}',f'{scale*float(x["VAN"]):.3f}',f'{scale*float(x["RD"]):.3f} ({scale*float(x["lower"]):.3f}, {scale*float(x["upper"]):.3f})'])
 table(d,f'Table S{num}. Remaining process summaries: '+('primary' if pop==P else 'full routine-culture'),['Measure','Unit','LZD','VAN','Difference (95% CI)'],dd,[2.1,.65,.78,.78,2.19],'Unconfirmed categories form an algebraic partition of single-threshold minus confirmed recovery, not a causal mediation decomposition. CIF area is a first-confirmation curve area, not actual days alive with sustained recovery. Tests and days are weighted means unless explicitly marked pooled. Testing includes observations after first recovery.',font=9)
gap=rows('reports/CONFIRMATION_GAP_v0.11.csv');table(d,'Table S10. Joint contrast between single-threshold and confirmed-recovery differences',['Population','LZD gap, pp','VAN gap, pp','Joint contrast (95% CI), pp'],[['Primary' if x['population']==P else 'Full routine-culture',f'{100*float(x["LZD"]):.2f}',f'{100*float(x["VAN"]):.2f}',f'{100*float(x["RD"]):.2f} ({100*float(x["lower"]):.2f}, {100*float(x["upper"]):.2f})'] for x in gap],[1.65,1.15,1.15,2.45],'This is the direct contrast of the two endpoint differences within each bootstrap draw.')
para(d,'S4. Additional sensitivity analyses','Heading 1')
para(d,'These post-primary analyses examined exposure route, missing baseline covariates, propensity-score behavior, treatment trajectories, and platelet observation opportunities.')

pq=rows('analysis_v0.14/output/PROPENSITY_WEIGHT_QUANTILES_v0.14.csv')
pqkeep=[]
for x in pq:
 if round(float(x['quantile']),2) in [.05,.50,.95,.99]:
  pqkeep.append([x['metric'].replace('_',' '),x['arm'],f"{100*float(x['quantile']):.0f}th",f"{float(x['value']):.4f}",f"{float(x['ESS']):.1f}"])
table(d,'Table S11. Propensity-score and overlap-weight distributions',['Measure','Arm','Quantile','Value','ESS'],pqkeep,[2.25,.65,.8,1.25,1.15],'Values are aggregate distribution summaries from the frozen primary propensity model. ESS is repeated within each arm for readability. Individual propensity scores and weights are not redistributed.')

rs=rows('analysis_v0.14/output/PARENTERAL_ROUTE_SENSITIVITY_v0.14.csv')[0]
mir=rows('analysis_v0.14/output/MULTIPLE_IMPUTATION_SENSITIVITY_v0.14.csv')
mp=next(x for x in mir if not x['imputation'])
sens=[
 ['Parenteral linezolid initiation',f"{100*float(rs['RD']):.2f} ({100*float(rs['lower']):.2f}, {100*float(rs['upper']):.2f})",f"{float(rs['ESS_LZD']):.1f} / {float(rs['ESS_VAN']):.1f}",f"{float(rs['max_abs_SMD']):.3f}",f"{rs['valid_draws']} / {rs['requested_draws']}"],
 ['Multiple imputation, pooled',f"{100*float(mp['RD']):.2f} ({100*float(mp['lower']):.2f}, {100*float(mp['upper']):.2f})",f"{float(mp['ESS_LZD']):.1f} / {float(mp['ESS_VAN']):.1f}",f"{float(mp['max_abs_SMD']):.3f}",f"{mp['bootstrap_valid']} / {mp['bootstrap_requested']}"]]
table(d,'Table S12. Route and missing-data sensitivity analyses',['Analysis','RD (95% CI), pp','ESS LZD / VAN','Maximum |SMD|','Valid / requested'],sens,[1.85,1.85,1.15,.9,.75],'The parenteral analysis used 2000 new patient-cluster refits. The multiple-imputation analysis used 20 completed datasets and 500 patient-cluster refits per dataset; variances were pooled with Rubin rules. Both were planned after the primary result was known and are exploratory.')

tr=rows('analysis_v0.14/output/TREATMENT_EXPOSURE_CONTEXT_v0.14.csv')
keys=[]
for x in tr:
 k=(x['metric'],x['unit'])
 if k not in keys:keys.append(k)
td=[]
for metric,unit in keys:
 z=[x for x in tr if x['metric']==metric and x['unit']==unit]
 td.append([metric,unit,next(x['value'] for x in z if x['arm']=='LZD'),next(x['value'] for x in z if x['arm']=='VAN')])
table(d,'Table S13. Observed treatment route and exposure trajectory',['Measure','Summary','LZD','VAN'],td,[3.25,1.0,1.0,1.0],'The observation window ended at day 14, hospital discharge or recorded death, whichever came first. Small and complementary initiation-route cells are suppressed. Values summarize recorded treatment during follow-up.')

ob=rows('analysis_v0.14/output/OBSERVATION_OPPORTUNITY_v0.14.csv')
om=[]
for metric in dict.fromkeys(x['metric'] for x in ob):
 z=[x for x in ob if x['metric']==metric]
 def fmt(arm):
  x=next(v for v in z if v['arm']==arm)
  if metric=='at_least_two_postinitiation_tests':return f"{100*float(x['weighted_mean']):.1f}%"
  return f"{float(x['weighted_mean']):.2f}; {float(x['weighted_median']):.2f} ({float(x['weighted_q25']):.2f}–{float(x['weighted_q75']):.2f})"
 om.append([metric.replace('_',' '),'Weighted mean; median (IQR)' if metric!='at_least_two_postinitiation_tests' else 'Weighted percentage',fmt('LZD'),fmt('VAN')])
table(d,'Table S14. Platelet observation opportunities',['Measure','Summary','LZD','VAN'],om,[2.6,1.7,1.1,1.1],'Testing was summarized through each hospital observation exit and includes measurements after first recovery.')

para(d,'Supplementary figure legends','Heading 1')
para(d,'Figure S1. Propensity-score, overlap-weight and baseline-balance diagnostics. Panels A and B show aggregate kernel-density curves for the frozen primary propensity scores and overlap weights. Panel C shows the 20 largest baseline absolute standardized mean differences among disclosure-safe fitted terms before and after weighting. The dashed line marks 0.10. Individual scores and weights are not displayed.')
para(d,'Figure S2. Platelet counts before live discharge and recorded monitoring. Panels A and B describe all live discharges within and after 14 days, including discharge after confirmed recovery. Panel C presents pooled tests per observation-day and the weighted mean of individual test rates. These summaries describe the recorded platelet observation process.')
d.save(O/'Additional_File_1_Supplement.docx')
# Concise reporting checklist with manuscript section locators.
d=normal_doc();para(d,'STROBE / RECORD reporting cross-reference','Title');para(d,'Author completion copy. Section locators are stable across pagination; ethics and author declarations remain pending. This is a completed mapping, not a claim that every unavailable source variable can be reconstructed.')
items=[('1','Title and abstract','Title; structured Abstract','Observational design and principal estimates'),('2–3','Background and objectives','Background','Prior recovery studies and incremental question'),('4–5','Design and setting','Methods: Design, data source, and analysis development','Single-center MIMIC-IV 3.1; exploratory local freeze'),('6','Participants','Methods: Population, exposure, and time zero; Figure 1','Cumulative filters; no future duration/survival condition'),('7–8','Variables and measurements','Methods: Baseline covariates; Outcomes; Supplement S1','Sources, code lists, time windows and ascertainment'),('9','Bias','Methods; Discussion limitations','Recording availability, indication, monitoring, competing events'),('10','Study size','Results: Cohort formation','All eligible admissions under frozen definitions; no post hoc power claim'),('11–12','Quantitative variables and statistics','Methods: Weighting; Statistical inference; Supplement','Transforms, missingness, cluster refits and sensitivity analyses'),('13','Participant flow','Figure 1; Results','Cumulative counts by treatment'),('14','Descriptive data','Table 1; Tables S2–S5','Before/after distributions and source missingness'),('15–16','Outcomes and main results','Table 2; Figures 2–3; Tables S6–S10','Raw event counts, weighted probabilities and uncertainty'),('17','Other analyses','Results: Sensitivity; Monitoring; Supplement S4','Post-result route, missing-data and observation-process analyses labeled exploratory'),('18–21','Interpretation and generalizability','Discussion; Conclusions','Sparse events, single center, residual bias, no causal mechanism'),('22','Funding','Declarations: Funding','Exact user-supplied funders; funder role not supplied'),('RECORD 1.1–1.3','Data type, source and linkage','Methods: Design and data source','Routinely collected records; internal supplied identifiers only'),('RECORD 6.1–6.3','Selection algorithms and validation','Supplement S1; source code; endpoint audit','Definitions and computational checks, no external chart-adjudication claim'),('RECORD 7.1','Code lists','Supplement S1; repository configuration','Frozen item, route, culture and medication definitions'),('RECORD 12.1–12.3','Access, cleaning and linkage','Methods; source code; Declarations','Internal linkage and availability checks; author access details pending'),('RECORD 13.1','Selection from database','Figure 1 legend; flow table','Pre-low-count initiation pool and cumulative selected cohort'),('RECORD 19.1','Routinely collected-data limitations','Discussion limitations','Misclassification, missingness, recording coverage and timing'),('RECORD 22.1','Supplemental protocols and code','Data availability; Additional files; GitHub','Frozen plans and aggregate reproducibility files')]
table(d,'Reporting items and locations',['Item','Topic','Manuscript location','Status / explanation'],items,[.67,1.35,2.2,2.28],font=9);d.save(O/'STROBE_RECORD_Checklist.docx')
# Author-facing copy/paste sheet (paragraph fields, not dense form grid).
abstract=resolved.split('## Abstract\n\n')[1].split('## Keywords')[0];abstract_plain=re.sub(r'^### ','',abstract,flags=re.M).strip();aw=len(re.findall(r'\b[\w−–]+\b',abstract_plain))
body=re.split(r'(?m)^## Background\n\n',resolved)[1].split('## List of abbreviations')[0];wc=len(re.findall(r'\b[\w−–]+\b',body))
d=normal_doc();d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.line_spacing=1.05;d.styles['Normal'].paragraph_format.space_after=Pt(4);d.styles['Heading 1'].paragraph_format.space_before=Pt(6);d.styles['Heading 1'].paragraph_format.space_after=Pt(3);para(d,'Submission copy / paste sheet','Title');para(d,'Target: BMC Pharmacology and Toxicology | Article type: Research article')
fields=[('Submission status','Local preparation. Do not submit until author contribution, ethics/access and all-author approval/exclusivity fields are completed.'),('Manuscript title',title),('Running title','Antibiotic initiation and platelet recovery'),('Keywords',resolved.split('## Keywords\n\n')[1].split('\n\n')[0]),('Authors in order','1. Kairui Liu (given: Kairui; family: Liu)\n2. Rong Chen (given: Rong; family: Chen)\n3. Yipei Huang (given: Yipei; family: Huang)\n4. Youxing Huang (given: Youxing; family: Huang; corresponding author)'),('Affiliation for all authors',AFFIL),('Corresponding author','Youxing Huang, MD\n'+ADDRESS+'\nTelephone: '+PHONE+'\nEmail: '+EMAIL),('Abstract',abstract_plain),('Word counts',f'Abstract including subheadings: {aw}; main body Background through Conclusions: {wc}. Recount in the submission system if its tokenizer differs.'),('Funding',fund),('Competing interests','The authors declare that they have no competing interests.'),('Consent for publication','Not applicable. No identifiable individual information is reported.'),('Data and code','MIMIC-IV 3.1: https://doi.org/10.13026/kpb9-mt58 — credentialed access required. Code: https://github.com/liukairui621/linezolid-platelet-recovery-mimiciv . No individual clinical records are redistributed.'),('AI assistance','Codex (OpenAI) and Claude Code (Anthropic) assisted with programming, documentation, manuscript preparation and technical review. Details are disclosed in Methods. These tools are not authors.'),('Required author completion','Actual CRediT contributions for each author; all-author final approval; originality and no simultaneous submission; local ethics approval/exemption/waiver and committee/reference if applicable; authorized MIMIC investigator/access information. Other author emails and ORCIDs have not been supplied. Do not invent them.'),('Files to upload','Main_Manuscript.docx (main article with tables and legends)\nTitle_Page.docx (only if requested separately)\nCover_Letter.docx (complete its pending declaration paragraph first)\nFigure_1_flow.pdf; Figure_2_first_events.pdf; Figure_3_recovery.pdf\nAdditional_File_1_Supplement.docx; Figure_S1_propensity_diagnostics.pdf; Figure_S2_observation_context.pdf\nAdditional_File_2_Aggregate_Data.zip\nSTROBE_RECORD_Checklist.docx'),('Cover-letter exclusivity statement — only after true','All authors have approved the final manuscript and agree to its submission. The manuscript has not been published previously and is not under consideration elsewhere. [Use only after all authors confirm.]'),('Suggested reviewers','Not supplied; no reviewers or contact details have been invented.'),('Submission portal','https://submission.nature.com/ — choose BMC Pharmacology and Toxicology. The author must perform the actual submission. Current author instructions are recorded in JOURNAL_SELECTION_AND_REQUIREMENTS.md.')]
for label,val in fields:
 para(d,label,'Heading 1');p=para(d,val)
 if label=='Abstract':p.paragraph_format.keep_together=True
d.save(O/'Submission_Copy_Paste_Sheet.docx');(O/'Submission_Copy_Paste_Sheet.txt').write_text('\n\n'.join(label+'\n'+val for label,val in fields),encoding='utf-8')
(O/'BUILD_CHECKS_v0.14.json').write_text(json.dumps({'build_utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),'runtime':'bundled primary runtime Python','abstract_words':aw,'body_words':wc,'reference_count':len(order),'files':[p.name for p in O.glob('*.docx')],'tables':'editable three-line, no shading or vertical/internal horizontal grid','author_declarations_pending':True,'manual_page_breaks_main':False,'line_and_page_numbering_main':True},indent=2),encoding='utf-8')
assert aw<=350
print(json.dumps({'docx':len(list(O.glob('*.docx'))),'abstract_words':aw,'body_words':wc,'references':len(order)}))
